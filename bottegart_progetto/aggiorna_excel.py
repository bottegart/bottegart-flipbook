# -*- coding: utf-8 -*-
"""Porta negli Excel il contenuto definitivo del flipbook.
Da eseguire dopo build_xlsx.py e make_flipbook.py:
- aggiunge al catalogo il foglio "Flipbook" (i libri del flipbook, con ISBN, trama, prezzo, età e fonti)
- rifà output/da_completare_per_editore.xlsx con la situazione aggiornata
"""
import json, collections
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

LIBRI = json.load(open('output/libri_flipbook.json', encoding='utf-8'))
WEB = json.load(open('dati_web.json', encoding='utf-8'))
FN_ = "Arial"
hfont = Font(name=FN_, bold=True, color="FFFFFF", size=10)
hfill = PatternFill("solid", start_color="1F3864")
bfont = Font(name=FN_, size=10)
thin = Side(style="thin", color="D9D9D9")
border = Border(top=thin, bottom=thin, left=thin, right=thin)
okfill = PatternFill("solid", start_color="E2EFDA")
nofill = PatternFill("solid", start_color="FCE4D6")

def tabella(ws, head, rows, widths, start=1):
    for c, h in enumerate(head, 1):
        cell = ws.cell(row=start, column=c, value=h)
        cell.font, cell.fill, cell.border = hfont, hfill, border
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for i, r in enumerate(rows, start + 1):
        for c, v in enumerate(r, 1):
            cell = ws.cell(row=i, column=c, value=v)
            cell.font, cell.border = bfont, border
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = ws.cell(row=start + 1, column=3)
    ws.auto_filter.ref = f"A{start}:{get_column_letter(len(head))}{start + len(rows)}"
    ws.row_dimensions[start].height = 30

def fonte(w):
    f = [w.get(k) for k in ("fonte", "fonte_prezzo", "fonte_eta", "fonte_sinossi") if w.get(k)]
    for r in w.get("sostituisci", []): f.append(r.get("fonte"))
    return " · ".join(dict.fromkeys(x for x in f if x))

# --- 1. foglio "Flipbook" nel catalogo
CAT = 'output/catalogo_case_editrici_indipendenti.xlsx'
wb = openpyxl.load_workbook(CAT)
if "Flipbook" in wb.sheetnames: del wb["Flipbook"]
ws = wb.create_sheet("Flipbook", 1)
HEAD = ["Fascia d’età", "Casa editrice", "Titolo", "Autori", "Età", "Prezzo copertina (€)", "ISBN", "CAA", "Copertina", "Collana", "Temi", "Sinossi", "Fonte dei dati presi dalla rete"]
rows = [[b["fascia"], b["editore"], b["titolo"], b["autore"], b["eta"], b["prezzo"], b["isbn"], "Sì" if b["caa"] else "", "Sì" if b["copertina"] else "No",
         b["collana"], b["temi"], b["sinossi"], fonte(b["web"])] for b in LIBRI]
tabella(ws, HEAD, rows, [20, 18, 40, 30, 16, 12, 16, 7, 10, 22, 26, 80, 50])
for i in range(2, len(rows) + 2): ws.cell(row=i, column=6).number_format = "#,##0.00"
wb.save(CAT)

# --- 2. da completare per editore
wb = Workbook()
ws = wb.active; ws.title = "Riepilogo"
ws["A1"] = "Situazione del flipbook per casa editrice"; ws["A1"].font = Font(name=FN_, bold=True, size=12)
cnt = collections.OrderedDict()
for b in LIBRI:
    c = cnt.setdefault(b["editore"], collections.Counter())
    c["libri"] += 1; c["copertina"] += b["copertina"]; c["trama"] += bool(b["sinossi"]); c["prezzo"] += b["prezzo"] is not None; c["eta"] += bool(b["eta"])
H = ["Casa editrice", "Libri nel flipbook", "Con copertina", "Con trama", "Con prezzo", "Con età", "Manca qualcosa?"]
R = [[p, c["libri"], c["copertina"], c["trama"], c["prezzo"], c["eta"], "No" if min(c["copertina"], c["trama"], c["prezzo"], c["eta"]) == c["libri"] else "Sì"] for p, c in cnt.items()]
tot = collections.Counter(); [tot.update(c) for c in cnt.values()]
R.append(["TOTALE", tot["libri"], tot["copertina"], tot["trama"], tot["prezzo"], tot["eta"], ""])
tabella(ws, H, R, [22, 14, 14, 12, 12, 12, 14], start=3)
for i, r in enumerate(R, 4):
    if r[6]: ws.cell(row=i, column=7).fill = okfill if r[6] == "No" else nofill
    if r[0] == "TOTALE":
        for c in range(1, 8): ws.cell(row=i, column=c).font = Font(name=FN_, bold=True, size=10)

ws2 = wb.create_sheet("Da completare")
manca = []
for b in LIBRI:
    m = [x for x, ok in (("copertina", b["copertina"]), ("trama", b["sinossi"]), ("prezzo", b["prezzo"] is not None), ("età", b["eta"])) if not ok]
    if m: manca.append([b["editore"], b["titolo"], b["autore"], b["isbn"], ", ".join(m)])
if not manca: manca = [["", "Non manca niente: tutti i libri del flipbook hanno copertina, trama, prezzo ed età.", "", "", ""]]
tabella(ws2, ["Casa editrice", "Titolo", "Autore", "ISBN", "Cosa manca"], manca, [20, 60, 30, 16, 24])

ws3 = wb.create_sheet("Modifiche")
mod = []
for k, w in WEB.items():
    p, t = k.split("|", 1)
    if w.get("rimuovi"): mod.append([p, t, "Tolto dal flipbook", w.get("motivo", "")])
    elif w.get("sostituisci"): mod.append([p, t, "Serie sostituita da titoli precisi", "; ".join(r["titolo"] for r in w["sostituisci"])])
    else:
        c = [x for x in ("autore_corretto", "isbn_corretto", "eta_corretta", "prezzo_corretto") if w.get(x)]
        if c: mod.append([p, t, "Corretto: " + ", ".join({"autore_corretto": "autore", "isbn_corretto": "ISBN", "eta_corretta": "età", "prezzo_corretto": "prezzo"}[x] for x in c), w.get("nota") or w.get("nota_prezzo") or ""])
tabella(ws3, ["Casa editrice", "Riga del catalogo", "Cosa è cambiato", "Dettagli"], mod, [20, 55, 34, 70])

ws4 = wb.create_sheet("Da controllare")
ws4["A1"] = ("Dati presi dalla rete (IBS e siti degli editori). Prima di mandare il link alle scuole conviene controllarne qualcuno a campione, "
             "e chiedere agli editori il permesso di usare testi e copertine.")
ws4["A1"].font = Font(name=FN_, italic=True, size=10); ws4["A1"].alignment = Alignment(wrap_text=True); ws4.row_dimensions[1].height = 30
chk = []
for b in LIBRI:
    w = b["web"]
    if not w or w.get("rimuovi"): continue
    cosa = [x for x, ks in (("trama", ("sinossi",)), ("età", ("eta", "eta_corretta", "fascia")), ("prezzo", ("prezzo", "prezzo_corretto")), ("ISBN", ("isbn", "isbn_corretto"))) if any(w.get(k) for k in ks)]
    if w.get("sostituisci"): cosa = ["titolo scelto dalla serie", "trama", "età", "prezzo", "ISBN"]
    if cosa: chk.append([b["editore"], b["titolo"], ", ".join(cosa), fonte(w)])
tabella(ws4, ["Casa editrice", "Titolo", "Dati presi dalla rete", "Fonte"], chk, [20, 55, 34, 70], start=3)
wb.save('output/da_completare_per_editore.xlsx')
print("Flipbook:", len(rows), "libri · da completare:", 0 if manca[0][0] == "" else len(manca), "· modifiche:", len(mod), "· da controllare:", len(chk))
