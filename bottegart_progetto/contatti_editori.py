# -*- coding: utf-8 -*-
"""Tabella dei contatti delle case editrici del catalogo e email per chiedere le condizioni di vendita.
Da eseguire dopo make_flipbook.py:  python contatti_editori.py  ->  output/contatti_editori.xlsx
Contatti verificati sui siti degli editori il 28 ottobre 2026 (colonna Fonte)."""
import json, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# casa editrice: (telefono, email, referente, note, fonte)
C = {
 "Uovonero": ("0373 500 622", "libri@uovonero.com", "",
              "Promozione Emme Promozione, distribuzione Messaggerie Libri.", "https://www.uovonero.com/"),
 "Terre di Mezzo": ("02 83242426", "segreteria@terre.it", "Stefania Falbo (librerie); Sabina Eleonori (scuole)",
              "Distribuzione Messaggerie Libri, promozione Emme Promozione. Le email dei referenti sono sulla pagina Contatti del sito.", "https://www.terre.it/contatti/"),
 "Edizioni Corsare": ("075 584 7055; WhatsApp 347 226 4702", "edizionicorsare@gmail.com", "",
              "Distribuzione Messaggerie Libri, promozione PDE.", "https://www.edizionicorsare.it/"),
 "Settenove": ("0721 787699 (8.30-13.30)", "ordini@settenove.it; amministrazione@settenove.it", "",
              "Promozione Emme Promozione, distribuzione Messaggerie Libri.", "https://www.settenove.it/contattaci"),
 "Camelozampa": ("0429 536477", "ordini@camelozampa.com; info@camelozampa.com", "",
              "Per le librerie di solito si passa da distributore (Messaggerie) o promotore (Emme Promozione).", "https://www.camelozampa.com/contatti/"),
 "BeccoGiallo": ("346 545 5244", "commerciale.beccogiallo@gmail.com; info@beccogiallo.it", "",
              "Acquisto diretto possibile per librerie, associazioni ed enti.", "https://www.beccogiallo.it/contatti/"),
 "Minibombo": ("0522 749356", "anna.biazzi@minibombo.it; info@minibombo.it", "Anna Biazzi (per i librai)",
              "Promozione PDE, distribuzione Messaggerie Libri.", "https://www.minibombo.it/contatti"),
 "L’Ippocampo": ("02 4851 3735; 377 255 1403", "r.magno@ippocampoedizioni.it; commerciale@ippocampoedizioni.it", "Ravi Magno (commerciale); ordini: Cristian Tettamanti",
              "Rapporto diretto con le librerie: gestiscono loro promozione e fatturazione. Ordini: ordini@ippocampoedizioni.it", "https://www.ippocampoedizioni.it/contatti/per-librai-e-rivenditori"),
 "Pelledoca": ("", "info@pelledocaeditore.it", "",
              "Nessun telefono sul sito. Distribuzione A.L.I. Agenzia Libraria Internazionale.", "https://pelledocaeditore.it/contatti/"),
 "Momo Edizioni": ("", "direzione@momoedizioni.it", "",
              "Nessun telefono sul sito.", "https://momoedizioni.it/"),
 "Lapis": ("06 3295 935", "commerciale.lapis@edizionilapis.it; scuola.lapis@edizionilapis.it", "",
              "Rete vendita PDE, distribuzione Messaggerie Libri.", "https://www.edizionilapis.it/contatti"),
 "La Meridiana": ("080 397 1945; cell. 329 839 1330", "info@lameridiana.it; media@lameridiana.it", "",
              "", "https://www.edizionilameridiana.it/contatti/"),
 "Dalia Edizioni": ("", "info@daliaedizioni.it; amministrazione@daliaedizioni.it (ordini)", "",
              "Solo email, nessun telefono.", "https://www.daliaedizioni.it/contatti/"),
 "Tunuè": ("0773 661760", "commerciale@tunue.com; info@tunue.com", "",
              "", "https://www.tunue.com/contatti/"),
 "Logos": ("059 412648", "commerciale@libri.it", "",
              "Casa editrice in chiusura: quantità e condizioni per lo stock da concordare con il commerciale.", "https://www.libri.it/contatti"),
 "effequ": ("06 9023 2210", "commerciale@effequ.it; librerie@effequ.it (incontri in libreria)", "",
              "Promozione LYR, distribuzione A.L.I.", "https://www.effequ.it/contatti/"),
 "Gallucci": ("06 8413 0332 (ordini); 06 841 3033", "biz@galluccieditore.com; info@galluccieditore.com", "",
              "Promozione PDE, distribuzione Messaggerie Libri.", "https://www.galluccieditore.com/it/contatti"),
 "Cronache Ribelli": ("WhatsApp ordini 348 175 0434", "ordini@cronacheribelli.it; info@cronacheribelli.it", "",
              "Sul sito: vendono alle librerie indipendenti solo in conto assoluto, minimo 5 libri; conto vendita non previsto.", "https://cronacheribelli.it/pages/distribuzione"),
 "People": ("0331 1629312", "amministrazione@peoplepub.it (rapporti con le librerie); ordini@peoplepub.it", "",
              "", "https://www.peoplepub.it/contatti"),
}

CATALOGO = "https://bottegart.github.io/bottegart-flipbook/"

def email(editore, n, referente):
    nome = referente.split(" (")[0].split(";")[0].strip() if referente else ""
    saluto = f"Gentile {nome}," if nome else f"Care amiche e cari amici di {editore},"
    libri = "un vostro titolo" if n == 1 else f"{n} vostri titoli"
    return f"""Oggetto: Microlibreria So’Stare – condizioni per librerie e conto vendita

{saluto}

vi scriviamo dall’Associazione Bottegart, che con molti di voi ha già avuto il piacere di lavorare. Per chi non ci conoscesse ancora: Bottegart unisce arte e diritti umani, e porta avanti la microlibreria So’Stare, un progetto sociale nato per sostenere il diritto universale allo studio e per portare buoni libri vicino a bambini, ragazzi e scuole.

Essendo un progetto sociale, sostenuto soprattutto dall’impegno delle persone che ci credono, la microlibreria è rimasta ferma per un periodo. Ora però ripartiamo con un rinnovato slancio: grazie alla campagna #ioleggoperché abbiamo ritrovato il contatto con tutte le scuole del nostro comprensorio, e la libreria è tornata a vivere.

Per le insegnanti abbiamo preparato un catalogo da sfogliare online, diviso per età, in cui abbiamo inserito {libri}: {CATALOGO}

Per organizzare al meglio gli ordini vi chiediamo:
- quali sono le vostre condizioni per una libreria in conto assoluto (sconto, eventuale minimo d’ordine, spese di spedizione e tempi di pagamento);
- se c’è la possibilità di lavorare anche in conto vendita, almeno per una parte dei titoli.

Ci farebbe molto piacere riprendere il dialogo con voi: restiamo a disposizione per una telefonata o qualsiasi informazione.

Un caro saluto e grazie,

[Nome e cognome]
Associazione Bottegart – Microlibreria So’Stare
bottegart@gmail.com
"""

L = json.load(open('output/libri_flipbook.json', encoding='utf-8'))
N = collections.Counter(b["editore"] for b in L)
assert set(N) <= set(C), set(N) - set(C)

FN_ = "Arial"
hfont = Font(name=FN_, bold=True, color="FFFFFF", size=10)
hfill = PatternFill("solid", start_color="1F3864")
bfont = Font(name=FN_, size=10)
thin = Side(style="thin", color="D9D9D9")
border = Border(top=thin, bottom=thin, left=thin, right=thin)
grey = PatternFill("solid", start_color="F2F2F2")

def tabella(ws, head, rows, widths):
    for c, h in enumerate(head, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font, cell.fill, cell.border = hfont, hfill, border
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for i, r in enumerate(rows, 2):
        for c, v in enumerate(r, 1):
            cell = ws.cell(row=i, column=c, value=v)
            cell.font, cell.border = bfont, border
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(head))}{len(rows) + 1}"
    ws.row_dimensions[1].height = 30

wb = Workbook()
ws = wb.active; ws.title = "Contatti"
ordine = sorted(N, key=lambda p: (-N[p], p))
rows = []
for p in ordine:
    tel, mail, ref, note, fonte = C[p]
    rows.append([p, N[p], tel or "non indicato sul sito", mail, ref or "—", note, "", "", fonte])
tabella(ws, ["Casa editrice", "Titoli nel catalogo", "Telefono", "Email", "Persona referente", "Note", "Contattata il", "Risposta / condizioni", "Fonte"],
        rows, [20, 10, 26, 44, 30, 46, 14, 40, 44])
for i in range(2, len(rows) + 2):
    for c in (7, 8): ws.cell(row=i, column=c).fill = grey

ws2 = wb.create_sheet("Email da inviare")
tabella(ws2, ["Casa editrice", "A chi scrivere", "Testo dell’email (pronto da copiare)"],
        [[p, C[p][1].split(";")[0].split(" (")[0], email(p, N[p], C[p][2])] for p in ordine], [20, 34, 110])
for i in range(2, len(ordine) + 2): ws2.row_dimensions[i].height = 420

ws3 = wb.create_sheet("Modello email")
ws3["A1"] = email("[casa editrice]", 0, "").replace("0 vostri titoli", "[numero] vostri titoli")
ws3["A1"].font = bfont; ws3["A1"].alignment = Alignment(wrap_text=True, vertical="top")
ws3.column_dimensions["A"].width = 120; ws3.row_dimensions[1].height = 520
wb.save('output/contatti_editori.xlsx')
print(len(rows), "case editrici")
